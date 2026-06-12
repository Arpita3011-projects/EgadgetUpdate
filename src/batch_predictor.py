"""
batch_predictor.py
──────────────────
Accepts a pandas DataFrame (from uploaded CSV/Excel in memory),
maps columns → features, runs predictions for every student,
returns rich results DataFrame + class summary + alerts.
"""

import os, sys
import numpy as np
import pandas as pd
import joblib

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from recommendation     import get_recommendations, calculate_addiction_score, get_alert_message
from form_column_mapper import (build_column_map, normalise_dataframe,
                                MODEL_FEATURES, IDENTITY_COLUMNS)

CLASS_NAMES = ['Low', 'Moderate', 'High', 'Severe']
RISK_COLORS = {'Low':'#27ae60','Moderate':'#f39c12','High':'#e67e22','Severe':'#c0392b'}


def load_artifacts(models_dir=None):
    if models_dir is None:
        models_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'models')
    model         = joblib.load(os.path.join(models_dir, 'random_forest_model.pkl'))
    scaler        = joblib.load(os.path.join(models_dir, 'scaler.pkl'))
    feature_names = joblib.load(os.path.join(models_dir, 'feature_names.pkl'))
    return model, scaler, feature_names


def read_uploaded_file(uploaded_file) -> pd.DataFrame:
    """
    Accept a Streamlit UploadedFile object (in-memory) OR a filepath string.
    Returns a DataFrame.
    """
    # Streamlit UploadedFile object
    if hasattr(uploaded_file, 'name'):
        name = uploaded_file.name.lower()
        if name.endswith('.csv'):
            return pd.read_csv(uploaded_file)
        elif name.endswith(('.xlsx', '.xls')):
            return pd.read_excel(uploaded_file)
        else:
            raise ValueError(f"Unsupported file: {uploaded_file.name}. Use .csv or .xlsx")
    # Filepath string
    elif isinstance(uploaded_file, str):
        ext = os.path.splitext(uploaded_file)[1].lower()
        if ext == '.csv':
            return pd.read_csv(uploaded_file)
        elif ext in ('.xlsx', '.xls'):
            return pd.read_excel(uploaded_file)
        else:
            raise ValueError(f"Unsupported file type: {ext}")
    else:
        raise TypeError(f"Expected file object or path string, got {type(uploaded_file)}")


def predict_from_dataframe(raw_df: pd.DataFrame, custom_column_map: dict = None) -> dict:
    """
    Core prediction engine — works directly with a DataFrame.

    Returns dict:
        results_df        : per-student predictions + features
        summary           : class-level stats dict
        alerts            : list of dicts for High/Severe students
        column_map        : {form_col: feature_name} mapping used
        missing_features  : features not found in input
    """
    # ── Column mapping ────────────────────────────────────────────────────
    column_map  = build_column_map(list(raw_df.columns), custom_column_map)
    features_df, identity_df = normalise_dataframe(raw_df, column_map)

    # ── Load model ────────────────────────────────────────────────────────
    model, scaler, feature_names = load_artifacts()

    missing = [f for f in feature_names if f not in features_df.columns]
    present = [f for f in feature_names if f in features_df.columns]

    if len(present) == 0:
        raise ValueError(
            "No recognisable model features found in the uploaded file.\n"
            "Make sure your CSV columns match the expected question texts, "
            "or check the Google Form Guide tab for the exact questions to use."
        )

    # Fill missing with neutral value
    for feat in missing:
        features_df[feat] = 0.5

    X        = features_df.reindex(columns=feature_names, fill_value=0).values
    X_scaled = scaler.transform(X)

    # ── Predict ───────────────────────────────────────────────────────────
    risk_indices = model.predict(X_scaled).astype(int)
    all_probs    = model.predict_proba(X_scaled)

    # ── Build results ─────────────────────────────────────────────────────
    results = identity_df.copy().reset_index(drop=True)

    results['predicted_risk']  = [CLASS_NAMES[i]              for i in risk_indices]
    results['risk_index']      = risk_indices
    results['addiction_score'] = [calculate_addiction_score(p) for p in all_probs]
    results['confidence_pct']  = [round(float(max(p))*100, 1)  for p in all_probs]
    results['prob_low']        = [round(p[0]*100, 1)            for p in all_probs]
    results['prob_moderate']   = [round(p[1]*100, 1)            for p in all_probs]
    results['prob_high']       = [round(p[2]*100, 1)            for p in all_probs]
    results['prob_severe']     = [round(p[3]*100, 1)            for p in all_probs]
    results['needs_alert']     = risk_indices >= 2

    for feat in feature_names:
        if feat in features_df.columns:
            results[feat] = features_df[feat].values

    results['top_recommendations'] = [
        " | ".join(get_recommendations(i)[:2]) for i in risk_indices
    ]

    # ── Summary ───────────────────────────────────────────────────────────
    total = len(results)
    dist  = results['predicted_risk'].value_counts()

    summary = {
        'total_students'     : total,
        'avg_addiction_score': round(results['addiction_score'].mean(), 1),
        'risk_distribution'  : {l: int(dist.get(l, 0)) for l in CLASS_NAMES},
        'risk_percentages'   : {l: round(int(dist.get(l,0))/total*100,1) for l in CLASS_NAMES},
        'high_severe_count'  : int(results['needs_alert'].sum()),
        'most_common_risk'   : results['predicted_risk'].mode()[0],
        'avg_screen_time'    : round(results['daily_screen_time_hours'].mean(),1)
                               if 'daily_screen_time_hours' in results else None,
        'avg_sleep_hours'    : round(results['sleep_hours'].mean(),1)
                               if 'sleep_hours' in results else None,
        'avg_gpa'            : round(results['gpa'].mean(),2)
                               if 'gpa' in results else None,
        'avg_stress'         : round(results['stress_level'].mean(),1)
                               if 'stress_level' in results else None,
    }

    if 'department' in results.columns:
        summary['department_avg_score'] = (
            results.groupby('department')['addiction_score']
            .mean().round(1).sort_values(ascending=False).to_dict()
        )

    # ── Alerts ────────────────────────────────────────────────────────────
    alerts = []
    for _, row in results[results['needs_alert']].iterrows():
        alerts.append({
            'name'          : row.get('student_name', f"Row {_+1}"),
            'student_id'    : row.get('student_id',  '—'),
            'department'    : row.get('department',  '—'),
            'risk'          : row['predicted_risk'],
            'score'         : row['addiction_score'],
            'risk_index'    : int(row['risk_index']),
            'alert_message' : get_alert_message(int(row['risk_index'])),
        })

    return {
        'results_df'      : results,
        'summary'         : summary,
        'alerts'          : alerts,
        'column_map'      : column_map,
        'missing_features': missing,
    }


def predict_batch(file_or_path, custom_column_map: dict = None) -> dict:
    """Convenience wrapper — accepts UploadedFile object or filepath string."""
    raw_df = read_uploaded_file(file_or_path)
    return predict_from_dataframe(raw_df, custom_column_map)


def export_results_to_bytes(results_df: pd.DataFrame) -> bytes:
    """Return colour-coded Excel as bytes (for Streamlit download button)."""
    import io
    from openpyxl.styles import PatternFill, Font

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        results_df.to_excel(writer, index=False, sheet_name='Results')
        ws = writer.sheets['Results']

        fill_map = {
            'Low'     : PatternFill('solid', fgColor='C8F7C5'),
            'Moderate': PatternFill('solid', fgColor='FFF3CD'),
            'High'    : PatternFill('solid', fgColor='FFD6A5'),
            'Severe'  : PatternFill('solid', fgColor='FFCDD2'),
        }
        headers = [c.value for c in ws[1]]
        try:
            risk_col = headers.index('predicted_risk') + 1
        except ValueError:
            risk_col = None

        for cell in ws[1]:
            cell.font = Font(bold=True)

        if risk_col:
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                risk_val = row[risk_col - 1].value
                if risk_val in fill_map:
                    for cell in row:
                        cell.fill = fill_map[risk_val]

        for col in ws.columns:
            max_len = max(len(str(c.value or '')) for c in col) + 2
            ws.column_dimensions[col[0].column_letter].width = min(max_len, 35)

    buf.seek(0)
    return buf.read()
